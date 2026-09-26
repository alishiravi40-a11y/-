"""E16 — Beta scheme bank statement (xls) vs Beta-system export (E15) vs Holoo 1404. AGGREGATES ONLY.

Usage: python3 E16_bank_statement_profile.py <bankStatements.xls> <beta_export.xlsx> --pg DSN
Names are used only in memory to link records; nothing personal is printed.
"""
import collections
import re
import statistics
import sys

import openpyxl
import psycopg
import xlrd

C = collections.Counter
FA = str.maketrans("۰۱۲۳۴۵۶۷۸۹٠١٢٣٤٥٦٧٨٩", "01234567890123456789")


def norm(s):
    return re.sub(r"\s+", " ", str(s).translate(FA).replace("ي", "ی").replace("ك", "ک").replace("‌", " ")).strip()


def key(s):
    return re.sub(r"[\s()]+", "", norm(re.sub(r"\(\s*بتا\s*\)|بتا\)|\(\s*بتا", "", str(s))))


def amt(x):
    return float(x) if isinstance(x, (int, float)) else float(str(x).translate(FA).replace(",", "").strip() or 0)


def main(bank, export, dsn):
    sh = xlrd.open_workbook(bank).sheets()[0]
    hdr = [norm(x) for x in sh.row_values(2)]
    B = [dict(zip(hdr, sh.row_values(r))) for r in range(3, sh.nrows) if str(sh.row_values(r)[hdr.index("تاریخ")]).strip()]
    for b in B:
        b["d"] = tuple(int(x) for x in re.sub(r"[^\d/]", "", norm(b["تاریخ"])).split("/"))
        b["dep"], b["wd"], b["id"], b["desc"] = amt(b["واریز"]), amt(b["برداشت"]), norm(b["شناسه واریز"]), norm(b["توضیحات"])
    inst = [b for b in B if b["dep"] > 0 and "قسط" in b["desc"]]
    for b in inst:
        b["name"] = key(re.sub(r"^.*?بابت\s*قسط", "", b["desc"]))
    bal = [amt(b["موجودی"]) for b in B]
    print("## E16.1 statement: account type from title:", "پس انداز" in norm(sh.row_values(1)[1]), "| rows", len(B),
          "| period", min(b["d"] for b in B), "-", max(b["d"] for b in B), "| running balance consistent (oldest-first / newest-first)",
          sum(abs(bal[i - 1] + B[i]["dep"] - B[i]["wd"] - bal[i]) < 1 for i in range(1, len(B))), "/",
          sum(abs(bal[i] + B[i - 1]["dep"] - B[i - 1]["wd"] - bal[i - 1]) < 1 for i in range(1, len(B))), "of", len(B) - 1)
    print("## E16.2 deposits: installment («بابت قسط») rows", len(inst), "sum", round(sum(b["dep"] for b in inst)),
          "| other deposits", sum(1 for b in B if b["dep"] > 0 and b not in inst), "sum", round(sum(b["dep"] for b in B if b["dep"] > 0 and b not in inst)),
          "| withdrawals", sum(b["wd"] > 0 for b in B), "sum", round(sum(b["wd"] for b in B)))
    day = C()
    for b in inst:
        day[f"{b['d'][1]:02}/{b['d'][2]:02}"] += round(b["dep"] / 1e6)
    print("## E16.3 installment deposits by day (million rials):", dict(sorted(day.items())))
    per = collections.defaultdict(list)
    for b in inst:
        per[b["id"]].append(b["dep"])
    print("## E16.4 deposit id: filled", sum(bool(b["id"]) for b in inst), "lengths", dict(C(len(b["id"]) for b in inst)),
          "| deposits per id", dict(C(map(len, per.values()))), "| ids with 2 deposits and equal amount",
          sum(1 for v in per.values() if len(v) == 2 and abs(v[0] - v[1]) < 1), "of", sum(1 for v in per.values() if len(v) == 2),
          "| one id -> one payer name", all(len({b["name"] for b in inst if b["id"] == i}) == 1 for i in per))
    wclass = C()
    for b in B:
        if b["wd"] > 0:
            t = b["desc"]
            wclass["bank fees/SMS/statement (< 1M)" if b["wd"] < 1e6 else
                   "instant PAYA to Refah (REFA)" if "REFA" in t else "supplier/other transfers (PAYA/SATNA)"] += 1
    refa = sum(b["wd"] for b in B if "REFA" in b["desc"])
    print("## E16.5 withdrawals by class:", dict(wclass), "| REFA total", round(refa), "=", round(100 * refa / sum(b["dep"] for b in inst), 2),
          "% of installment deposits | any withdrawal described as bank share/installment fee:",
          sum(1 for b in B if b["wd"] > 0 and re.search("قسط|سهم|بتا", b["desc"])))
    ws = openpyxl.load_workbook(export, data_only=True).active
    H = [norm(ws.cell(1, c).value or "") for c in range(1, ws.max_column + 1)]
    X = [{H[c - 1]: ws.cell(r, c).value for c in range(1, ws.max_column + 1)} for r in range(2, ws.max_row + 1) if ws.cell(r, 1).value]
    xnames = {key((x["نام مشتری"] or "") + (x["نام خانوادگی مشتری"] or "")) for x in X}
    xnc = {str(x["کدملی مشتری"]).strip().zfill(10) for x in X}
    with psycopg.connect(dsn) as c:
        hb = c.execute("""SELECT name FROM holoo_mirror.account WHERE source_db='holoo1_1404' AND length(code)>=7
                          AND (name ~ '\\(\\s*بتا|بتا\\s*\\)' OR left(code,7)='1080004') AND code<>'1080004'""").fetchall()
        persons = c.execute("SELECT name, national_code FROM holoo_mirror.person WHERE source_db='holoo1_1404'").fetchall()
        coll = c.execute("""WITH b AS (SELECT code, name FROM holoo_mirror.account WHERE source_db='holoo1_1404' AND length(code)>=7
                          AND (name ~ '\\(\\s*بتا|بتا\\s*\\)' OR left(code,7)='1080004') AND code<>'1080004')
                          SELECT b.name, l.credit FROM holoo_mirror.voucher_line l JOIN b ON b.code=l.account_code
                          WHERE l.source_db='holoo1_1404' AND l.in_ledger AND l.credit>0 AND l.description LIKE '%وصول قسط%'""").fetchall()
    hbk = {key(n) for (n,) in hb}
    pk = collections.defaultdict(set)
    for n, nc in persons:
        pk[key(n)].add(nc)
    names = {b["name"] for b in inst}
    print("## E16.6 payers", len(names), "| = Holoo 1404 beta customers", len(names & hbk), "| = Beta-export customers (name)", len(names & xnames),
          "| deposit id = Beta-export national code", sum(b["id"] in xnc for b in inst), "of", len(inst),
          "| deposit id = national code of same-named Holoo person", sum(1 for b in inst if b["name"] in pk and b["id"] in pk[b["name"]]),
          "of", sum(1 for b in inst if b["name"] in pk))
    h = collections.defaultdict(list)
    for n, cr in coll:
        h[key(n)].append(round(cr))
    cmp = C()
    for n in names:
        if n in h:
            d = statistics.mode([round(b["dep"]) for b in inst if b["name"] == n])
            cmp["equal" if abs(statistics.mode(h[n]) - d) < 2 else "different"] += 1
    print("## E16.7 1405 deposit vs modal 1404 Holoo «وصول قسط» credit of the same customer:", dict(cmp))


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2], sys.argv[sys.argv.index("--pg") + 1])
