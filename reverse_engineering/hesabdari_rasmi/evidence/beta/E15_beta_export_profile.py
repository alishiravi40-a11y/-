"""E15 — Profile of a Beta-system contracts export (Excel) for INFO-1. Prints AGGREGATES ONLY (no personal values).

Usage: python3 E15_beta_export_profile.py <export.xlsx> [--pg DSN]   (--pg: overlap with Holoo persons by national code)
"""
import collections
import re
import sys

import jdatetime
import openpyxl

C = collections.Counter
PERSONAL = {"نام مشتری", "نام خانوادگی مشتری", "شماره موبایل", "تاریخ تولد مشتری", "ثبت کننده"}   # never printed


def jd(s):
    return tuple(int(x) for x in str(s).split("/"))


def eom(y, m):
    return 31 if m <= 6 else (30 if m <= 11 else (30 if jdatetime.date(y, 1, 1).isleap() else 29))


def num(x):
    if isinstance(x, (int, float)):
        return x
    return float(str(x).replace(",", "").translate(str.maketrans("۰۱۲۳۴۵۶۷۸۹", "0123456789")).strip() or 0)


def valid_nc(s):
    s = s.zfill(10)
    if not re.fullmatch(r"\d{10}", s) or len(set(s)) == 1:
        return False
    t = sum(int(s[i]) * (10 - i) for i in range(9)) % 11
    return int(s[9]) == (t if t < 2 else 11 - t)


def due_count(r, E):
    y, m, _ = jd(r["تاریخ اولین قسط"])
    return sum(1 for i in range(r["اقساط"])
               if (y + (m - 1 + i) // 12, (m - 1 + i) % 12 + 1, eom(y + (m - 1 + i) // 12, (m - 1 + i) % 12 + 1)) <= E)


def main(path, dsn=None):
    ws = openpyxl.load_workbook(path, data_only=True).active
    H = [(ws.cell(1, c).value or "").strip() for c in range(1, ws.max_column + 1)]
    rows = [{H[c - 1]: ws.cell(r, c).value for c in range(1, ws.max_column + 1)} for r in range(2, ws.max_row + 1) if ws.cell(r, 1).value]
    print("## E15.1 shape: rows", len(rows), "columns", len(H), "| headers:", [h for h in H if h not in PERSONAL])
    ids = [r["شناسه"] for r in rows]
    order = sorted(rows, key=lambda r: (jd(r["تاریخ ثبت"]), str(r["زمان ثبت"])))
    print("## E15.2 شناسه: unique", len(set(ids)) == len(ids), "| increasing with registration time",
          sum(b["شناسه"] > a["شناسه"] for a, b in zip(order, order[1:])), "of", len(order) - 1,
          "| registration range", min(jd(r["تاریخ ثبت"]) for r in rows), "-", max(jd(r["تاریخ ثبت"]) for r in rows))
    nc = [str(r["کدملی مشتری"]).strip().zfill(10) for r in rows]
    acc = [str(r["شماره حساب"]).strip() for r in rows]
    per = C(nc)
    pa = collections.defaultdict(set); ap = collections.defaultdict(set)
    for n, a in zip(nc, acc):
        pa[n].add(a); ap[a].add(n)
    print("## E15.3 national code: valid checksum", sum(map(valid_nc, nc)), "of", len(nc), "| customers", len(per),
          "| contracts per customer", dict(C(per.values())))
    print("## E15.4 شماره حساب: length", dict(C(map(len, acc))), "| distinct", len(set(acc)), "| accounts per customer",
          dict(C(len(v) for v in pa.values())), "| customers per account", dict(C(len(v) for v in ap.values())))
    tot = [num(r["مبلغ کل اعتباری"]) for r in rows]; first = [num(r["مبلغ اولین قسط"]) for r in rows]; n = [r["اقساط"] for r in rows]
    print("## E15.5 installments: count distribution", dict(sorted(C(n).items())),
          "| first == floor(total/count)", sum(f == t // k for t, f, k in zip(tot, first, n)),
          "| remainder (total - first*count) in [0, count)", sum(0 <= t - f * k < k for t, f, k in zip(tot, first, n)),
          "| remainder == 0", sum(t == f * k for t, f, k in zip(tot, first, n)))
    fd = [jd(r["تاریخ اولین قسط"]) for r in rows]
    x = collections.defaultdict(C)
    for r, d in zip(rows, fd):
        a = jd(r["تاریخ ثبت"]); x[(a[2] <= 15)][(d[0] - a[0]) * 12 + (d[1] - a[1])] += 1
    print("## E15.6 first due: end of Jalali month", sum(d[2] == eom(d[0], d[1]) for d in fd), "of", len(fd),
          "| months after registration: reg day<=15", dict(x[True]), "reg day>15", dict(x[False]))
    col = [r["اقساط وصولی"] for r in rows]; due = [r["اقساط سررسید"] for r in rows]; dl = [r["اقساط حذف شده"] for r in rows]
    camt = [num(r["مبلغ اقساط وصولی"]) for r in rows]
    live = [i for i in range(len(rows)) if dl[i] == 0]
    E = max(jd(r["تاریخ ثبت"]) for r in rows)
    print("## E15.7 collections: collected_amount == collected*first", sum(abs(a - c * f) < 1 for a, c, f in zip(camt, col, first)),
          "| live contracts", len(live), "| collected + «اقساط سررسید» == installments due by export date",
          sum(col[i] + due[i] == due_count(rows[i], E) for i in live), "| live with overdue>0", sum(due[i] > 0 for i in live),
          "| fully collected", sum(col[i] == n[i] for i in live))
    print("## E15.8 deleted installments: rows with deleted>0", sum(d > 0 for d in dl), "| whole contract deleted", sum(d == k for d, k in zip(dl, n)),
          "| partial (count, deleted, collected, overdue)", [(n[i], dl[i], col[i], due[i]) for i in range(len(rows)) if 0 < dl[i] < n[i]],
          "| deleted + collected == count", sum(dl[i] + col[i] == n[i] for i in range(len(rows)) if dl[i] > 0), "of", sum(d > 0 for d in dl))
    print("## E15.9 مبلغ پرداخت نقدی values", C(num(r["مبلغ پرداخت نقدی"]) for r in rows).most_common(5),
          "| distinct: سازمان", len({r["سازمان مشتری"] for r in rows}), "بازاریاب", len({r["بازاریاب"] for r in rows}),
          "نمایندگی empty", sum(not r["نمایندگی"] for r in rows), "registrars", len({r["کدملی ثبت کننده"] for r in rows}),
          "| عنوان generic", sum(str(r["عنوان"]).strip() == "بابت خرید کالا" for r in rows))
    if dsn:
        import psycopg
        with psycopg.connect(dsn) as c:
            c.execute("CREATE TEMP TABLE x (nc text, reg text)")
            with c.cursor().copy("COPY x FROM STDIN") as cp:
                for r, k in zip(rows, nc):
                    cp.write_row((k, str(r["تاریخ ثبت"])))
            print("## E15.10 overlap with Holoo 1404 persons (customers, in_holoo, as_beta_in_holoo, registered_in_1404, of_which_beta_in_holoo):",
                  c.execute("""WITH b AS (SELECT code FROM holoo_mirror.account WHERE source_db='holoo1_1404' AND length(code)>=7
                                 AND (name ~ '\\(\\s*بتا|بتا\\s*\\)' OR left(code,7)='1080004') AND code<>'1080004'),
                    p AS (SELECT national_code n, bool_or(debit_account IN (SELECT code FROM b)) beta FROM holoo_mirror.person
                          WHERE source_db='holoo1_1404' AND national_code ~ '^[0-9]{10}$' GROUP BY 1)
                    SELECT count(DISTINCT nc), count(DISTINCT nc) FILTER (WHERE p.n IS NOT NULL), count(DISTINCT nc) FILTER (WHERE p.beta),
                           count(DISTINCT nc) FILTER (WHERE reg < '1405'), count(DISTINCT nc) FILTER (WHERE reg < '1405' AND p.beta)
                    FROM x LEFT JOIN p ON p.n = x.nc""").fetchone())


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[sys.argv.index("--pg") + 1] if "--pg" in sys.argv else None)
